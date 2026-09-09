import io
import csv
from datetime import datetime, date
from flask import Blueprint, render_template, request, send_file, flash, redirect, url_for
from flask_login import login_required, current_user
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from models import db
from models.transaction import Transaction
from models.category import Category
from models.account import Account
from ai.insights import FinancialInsightsEngine

reports_bp = Blueprint('reports', __name__)


def get_filtered_transactions(user_id, start_date=None, end_date=None, category_id=None, account_id=None, tx_type=None):
    query = Transaction.query.filter_by(user_id=user_id)
    if start_date:
        query = query.filter(Transaction.transaction_date >= start_date)
    if end_date:
        query = query.filter(Transaction.transaction_date <= end_date)
    if category_id:
        query = query.filter(Transaction.category_id == category_id)
    if account_id:
        query = query.filter(Transaction.account_id == account_id)
    if tx_type in ['expense', 'income']:
        query = query.filter(Transaction.transaction_type == tx_type)
    return query.order_by(Transaction.transaction_date.desc()).all()


@reports_bp.route('/')
@login_required
def index():
    categories = Category.query.filter_by(user_id=current_user.id).all()
    accounts = Account.query.filter_by(user_id=current_user.id).all()
    today = date.today()
    
    monthly_report = FinancialInsightsEngine.generate_monthly_report(current_user.id, today.year, today.month)
    health_score = FinancialInsightsEngine.calculate_financial_health_score(current_user.id)

    return render_template(
        'reports/index.html',
        categories=categories,
        accounts=accounts,
        monthly_report=monthly_report,
        health_score=health_score,
        today=today
    )


@reports_bp.route('/export/csv')
@login_required
def export_csv():
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    cat_id = request.args.get('category_id', type=int)
    acc_id = request.args.get('account_id', type=int)
    tx_type = request.args.get('type')

    start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else None
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date() if end_date_str else None

    txs = get_filtered_transactions(current_user.id, start_date, end_date, cat_id, acc_id, tx_type)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(['ID', 'Date', 'Type', 'Amount', 'Category', 'Account', 'Description', 'Is Recurring', 'Is Anomaly'])

    for t in txs:
        cat_name = t.category.category_name if t.category else 'Uncategorized'
        acc_name = t.account.account_name if t.account else 'Unknown'
        writer.writerow([
            t.id,
            t.transaction_date.strftime('%Y-%m-%d'),
            t.transaction_type.capitalize(),
            f"{float(t.amount):.2f}",
            cat_name,
            acc_name,
            t.description or '',
            'Yes' if t.is_recurring else 'No',
            'Yes' if t.is_anomaly else 'No'
        ])

    output.seek(0)
    filename = f"financial_report_{date.today().strftime('%Y%m%d')}.csv"
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=filename
    )


@reports_bp.route('/export/excel')
@login_required
def export_excel():
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    cat_id = request.args.get('category_id', type=int)
    acc_id = request.args.get('account_id', type=int)
    tx_type = request.args.get('type')

    start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else None
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date() if end_date_str else None

    txs = get_filtered_transactions(current_user.id, start_date, end_date, cat_id, acc_id, tx_type)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions Report"

    # Style definitions
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = ['ID', 'Date', 'Type', 'Amount ($)', 'Category', 'Account', 'Description', 'Recurring', 'Anomaly']
    ws.append(headers)

    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_income = 0.0
    total_expense = 0.0

    for t in txs:
        cat_name = t.category.category_name if t.category else 'Uncategorized'
        acc_name = t.account.account_name if t.account else 'Unknown'
        amt = float(t.amount)
        if t.transaction_type == 'income':
            total_income += amt
        else:
            total_expense += amt

        row = [
            t.id,
            t.transaction_date.strftime('%Y-%m-%d'),
            t.transaction_type.capitalize(),
            amt,
            cat_name,
            acc_name,
            t.description or '',
            'Yes' if t.is_recurring else 'No',
            'Yes' if t.is_anomaly else 'No'
        ]
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.font = data_font
            cell.border = thin_border

    # Summary Rows
    ws.append([])
    ws.append(['', '', 'Total Income:', total_income, '', '', '', '', ''])
    ws.append(['', '', 'Total Expenses:', total_expense, '', '', '', '', ''])
    ws.append(['', '', 'Net Savings:', total_income - total_expense, '', '', '', '', ''])

    # Set column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"financial_report_{date.today().strftime('%Y%m%d')}.xlsx"

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )


@reports_bp.route('/export/pdf')
@login_required
def export_pdf():
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    cat_id = request.args.get('category_id', type=int)
    acc_id = request.args.get('account_id', type=int)
    tx_type = request.args.get('type')

    start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date() if start_date_str else None
    end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date() if end_date_str else None

    txs = get_filtered_transactions(current_user.id, start_date, end_date, cat_id, acc_id, tx_type)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    elements = []
    styles = getSampleStyleSheet()

    # Title & Subtitle
    title_style = ParagraphStyle(
        'ReportTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1e293b'),
        alignment=0
    )
    elements.append(Paragraph("<b>Smart Expense Tracker — Financial Statement</b>", title_style))
    elements.append(Paragraph(f"Generated for <b>{current_user.name}</b> ({current_user.email}) on {date.today().strftime('%B %d, %Y')}", styles['Normal']))
    elements.append(Spacer(1, 14))

    # Summary Box
    total_income = sum(float(t.amount) for t in txs if t.transaction_type == 'income')
    total_expense = sum(float(t.amount) for t in txs if t.transaction_type == 'expense')
    net_savings = total_income - total_expense

    summary_data = [
        ['Total Income', 'Total Expenses', 'Net Balance / Savings', 'Total Transactions'],
        [f"${total_income:,.2f}", f"${total_expense:,.2f}", f"${net_savings:,.2f}", str(len(txs))]
    ]
    summary_table = Table(summary_data, colWidths=[130, 130, 150, 130])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f46e5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#f8fafc')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 18))

    # Transactions Table
    elements.append(Paragraph("<b>Recent Transaction Details</b>", styles['Heading2']))
    elements.append(Spacer(1, 6))

    table_data = [['Date', 'Type', 'Category', 'Account', 'Description', 'Amount']]
    for t in txs[:60]:  # Limit to 60 for clean PDF layout
        cat_name = t.category.category_name if t.category else 'Uncategorized'
        acc_name = t.account.account_name if t.account else 'Unknown'
        amt_str = f"+${float(t.amount):,.2f}" if t.transaction_type == 'income' else f"-${float(t.amount):,.2f}"
        table_data.append([
            t.transaction_date.strftime('%Y-%m-%d'),
            t.transaction_type.capitalize(),
            cat_name[:18],
            acc_name[:15],
            (t.description or '')[:25],
            amt_str
        ])

    tx_table = Table(table_data, colWidths=[65, 55, 105, 95, 140, 80])
    tx_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e293b')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(tx_table)

    doc.build(elements)
    buffer.seek(0)
    filename = f"financial_report_{date.today().strftime('%Y%m%d')}.pdf"

    return send_file(
        buffer,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=filename
    )
